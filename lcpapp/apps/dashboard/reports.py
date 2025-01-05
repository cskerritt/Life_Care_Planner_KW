from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from django.utils.translation import gettext as _

class CarePlanReportGenerator:
    def __init__(self, care_plan, sections):
        self.care_plan = care_plan
        self.sections = sections
        self.styles = getSampleStyleSheet()
        
        # Custom styles
        self.styles.add(ParagraphStyle(
            name='CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=24,
            spaceAfter=30
        ))
        
        self.styles.add(ParagraphStyle(
            name='SectionTitle',
            parent=self.styles['Heading2'],
            fontSize=16,
            spaceAfter=20
        ))
    
    def generate_pdf(self, response):
        doc = SimpleDocTemplate(response, pagesize=letter)
        story = []
        
        # Title
        title = Paragraph(
            f"Life Care Plan Report - {self.care_plan.evaluee.full_name}",
            self.styles['CustomTitle']
        )
        story.append(title)
        story.append(Spacer(1, 12))
        
        # Add requested sections
        if 'summary' in self.sections:
            story.extend(self._generate_summary_section())
        
        if 'evaluee' in self.sections:
            story.extend(self._generate_evaluee_section())
        
        if 'equipment' in self.sections:
            story.extend(self._generate_equipment_section())
        
        if 'home_care' in self.sections:
            story.extend(self._generate_home_care_section())
        
        if 'costs' in self.sections:
            story.extend(self._generate_costs_section())
        
        doc.build(story)
    
    def generate_excel(self, response):
        wb = Workbook()
        
        # Create sheets for each section
        if 'summary' in self.sections:
            self._generate_summary_sheet(wb)
        
        if 'evaluee' in self.sections:
            self._generate_evaluee_sheet(wb)
        
        if 'equipment' in self.sections:
            self._generate_equipment_sheet(wb)
        
        if 'home_care' in self.sections:
            self._generate_home_care_sheet(wb)
        
        if 'costs' in self.sections:
            self._generate_costs_sheet(wb)
        
        # Remove default sheet if it exists
        if 'Sheet' in wb.sheetnames:
            wb.remove(wb['Sheet'])
        
        wb.save(response)
    
    def _generate_summary_section(self):
        story = []
        story.append(Paragraph(_("Executive Summary"), self.styles['SectionTitle']))
        
        # Add summary content
        summary_text = f"""
        This Life Care Plan has been prepared for {self.care_plan.evaluee.full_name}.
        Date of Birth: {self.care_plan.evaluee.date_of_birth}
        Date of Evaluation: {self.care_plan.created_at}
        """
        story.append(Paragraph(summary_text, self.styles['Normal']))
        story.append(Spacer(1, 12))
        
        return story
    
    def _generate_evaluee_section(self):
        story = []
        story.append(Paragraph(_("Evaluee Information"), self.styles['SectionTitle']))
        
        # Create evaluee info table
        data = [
            [_("Name"), self.care_plan.evaluee.full_name],
            [_("Date of Birth"), self.care_plan.evaluee.date_of_birth],
            [_("Gender"), self.care_plan.evaluee.gender],
            [_("Contact"), self.care_plan.evaluee.contact_info],
        ]
        
        t = Table(data, colWidths=[2*inch, 4*inch])
        t.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 1, colors.black),
            ('BACKGROUND', (0,0), (0,-1), colors.lightgrey),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        
        story.append(t)
        story.append(Spacer(1, 12))
        
        return story
    
    def _generate_equipment_section(self):
        story = []
        story.append(Paragraph(_("Equipment & Supplies"), self.styles['SectionTitle']))
        
        # Create equipment table
        headers = [_("Item"), _("Quantity"), _("Frequency"), _("Cost")]
        data = [headers]
        
        for equipment in self.care_plan.equipment_set.all():
            data.append([
                equipment.name,
                equipment.quantity,
                equipment.get_frequency_display(),
                f"${equipment.cost:.2f}"
            ])
        
        t = Table(data, colWidths=[3*inch, 1*inch, 1.5*inch, 1.5*inch])
        t.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 1, colors.black),
            ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        
        story.append(t)
        story.append(Spacer(1, 12))
        
        return story
    
    def _generate_home_care_section(self):
        story = []
        story.append(Paragraph(_("Home Care Services"), self.styles['SectionTitle']))
        
        # Create home care table
        headers = [_("Service"), _("Provider"), _("Hours/Visit"), _("Cost/Hour")]
        data = [headers]
        
        for service in self.care_plan.homecare_set.all():
            data.append([
                service.get_service_type_display(),
                service.provider_name,
                service.hours_per_visit,
                f"${service.rate_per_hour:.2f}"
            ])
        
        t = Table(data, colWidths=[2*inch, 2*inch, 1.5*inch, 1.5*inch])
        t.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 1, colors.black),
            ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        
        story.append(t)
        story.append(Spacer(1, 12))
        
        return story
    
    def _generate_costs_section(self):
        story = []
        story.append(Paragraph(_("Cost Analysis"), self.styles['SectionTitle']))
        
        # Calculate total costs
        equipment_costs = sum(e.annual_cost for e in self.care_plan.equipment_set.all())
        home_care_costs = sum(h.annual_cost for h in self.care_plan.homecare_set.all())
        total_costs = equipment_costs + home_care_costs
        
        # Create costs table
        data = [
            [_("Category"), _("Annual Cost")],
            [_("Equipment & Supplies"), f"${equipment_costs:.2f}"],
            [_("Home Care Services"), f"${home_care_costs:.2f}"],
            [_("Total Annual Cost"), f"${total_costs:.2f}"],
        ]
        
        t = Table(data, colWidths=[4*inch, 2*inch])
        t.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 1, colors.black),
            ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
            ('BACKGROUND', (0,-1), (-1,-1), colors.lightgrey),
            ('PADDING', (0,0), (-1,-1), 6),
        ]))
        
        story.append(t)
        story.append(Spacer(1, 12))
        
        return story
    
    def _generate_summary_sheet(self, wb):
        ws = wb.create_sheet(_("Summary"))
        ws.append([_("Life Care Plan Summary")])
        ws.append([])
        ws.append([_("Evaluee Name"), self.care_plan.evaluee.full_name])
        ws.append([_("Date of Birth"), self.care_plan.evaluee.date_of_birth])
        ws.append([_("Date of Evaluation"), self.care_plan.created_at])
    
    def _generate_evaluee_sheet(self, wb):
        ws = wb.create_sheet(_("Evaluee Info"))
        ws.append([_("Evaluee Information")])
        ws.append([])
        ws.append([_("Name"), self.care_plan.evaluee.full_name])
        ws.append([_("Date of Birth"), self.care_plan.evaluee.date_of_birth])
        ws.append([_("Gender"), self.care_plan.evaluee.gender])
        ws.append([_("Contact"), self.care_plan.evaluee.contact_info])
    
    def _generate_equipment_sheet(self, wb):
        ws = wb.create_sheet(_("Equipment"))
        ws.append([_("Equipment & Supplies")])
        ws.append([])
        ws.append([_("Item"), _("Quantity"), _("Frequency"), _("Cost")])
        
        for equipment in self.care_plan.equipment_set.all():
            ws.append([
                equipment.name,
                equipment.quantity,
                equipment.get_frequency_display(),
                equipment.cost
            ])
    
    def _generate_home_care_sheet(self, wb):
        ws = wb.create_sheet(_("Home Care"))
        ws.append([_("Home Care Services")])
        ws.append([])
        ws.append([_("Service"), _("Provider"), _("Hours/Visit"), _("Cost/Hour")])
        
        for service in self.care_plan.homecare_set.all():
            ws.append([
                service.get_service_type_display(),
                service.provider_name,
                service.hours_per_visit,
                service.rate_per_hour
            ])
    
    def _generate_costs_sheet(self, wb):
        ws = wb.create_sheet(_("Costs"))
        ws.append([_("Cost Analysis")])
        ws.append([])
        
        equipment_costs = sum(e.annual_cost for e in self.care_plan.equipment_set.all())
        home_care_costs = sum(h.annual_cost for h in self.care_plan.homecare_set.all())
        total_costs = equipment_costs + home_care_costs
        
        ws.append([_("Category"), _("Annual Cost")])
        ws.append([_("Equipment & Supplies"), equipment_costs])
        ws.append([_("Home Care Services"), home_care_costs])
        ws.append([_("Total Annual Cost"), total_costs])
